#!/usr/bin/env python3
"""
Regression tests for the founder-vs-market model, the workbook, the recalculator and the report builder.

    python scripts/selftest.py [--recalc /path/to/recalc.py] [--python-only]

Known-answer checks on the Python model; input-validation checks; report-builder checks (placeholders, number
verification, required content, chart manifests, rendered pagination); then a workbook per revenue unit plus
multi-line workbooks (including colliding and case-only-different line names), each recalculated with LibreOffice
via scripts/recalc.py and compared with Python for all four cases and the Combined sheet.
The Excel comparison is MANDATORY: without a working recalculation the run exits 1 with "INCOMPLETE".
--python-only skips Excel for local iteration (never a release-grade run).
"""
import contextlib, io, json, os, shutil, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from financial_model import run_all, run_line, validate, CASES  # noqa: E402
import financial_model as fm  # noqa: E402


def _raises(fn, needle):
    try: fn()
    except Exception as e: return needle.lower() in str(e).lower()
    return False


FAILS = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else f"  — {detail}"))
    if not cond: FAILS.append(name)


def close(a, b, tol=0.02, abs_tol=0.0):
    if a is None or b is None: return a == b
    if isinstance(a, str) or isinstance(b, str): return str(a) == str(b)
    return abs(a - b) <= max(tol * max(1.0, abs(b)), abs_tol)


def A(founder=None, market=None, status="founder_stated", source="test source"):
    d = {"status": status} if founder is not None else {}
    if founder is not None: d["founder"] = founder
    if market is not None: d["market"] = market; d["source"] = source
    return d


FUNNEL = [{"step": "signups / month", "value": "500", "source": "founder_stated"}, {"step": "trial→paid", "value": "4–15 %", "source": "conv.trial.nocard.2026"}]
G = lambda **kw: dict({"motion": "organic", "ramp_k": 0, "funnel": FUNNEL}, **kw)
FP = {"funding_path": "pre_seed", "currency": "USD"}

CASES_IN = {
    # q = 1 edge case: churn 20 %, expansion 25 %
    "subscription": {"revenue_unit": "subscription", "growth": G(), "fte": 0,
                     "assumptions": {"pricing": A(100, 100), "gross_margin": A(80, 80), "monthly_churn": A(20, 20), "expansion": A(25, 25), "cac": A(120, 120),
                                     "new_customers_per_month": A(1, 1), "fte_cost": A(0, 0)}},
    "subscription_growth": {"revenue_unit": "subscription", "growth": G(motion="plg", ramp_k=1.0, sales_cycle_months=1), "fte": 0,
                            "comparable_benchmark": {"segment_size": 5000, "comparable_penetration_pct": [2, 5]},
                            "assumptions": {"pricing": A(70, [50, 80]), "gross_margin": A(80, 80), "monthly_churn": A(3, [3, 6], "founder_guess"), "expansion": A(1, [0.5, 2]),
                                            "cac": A(150, [150, 300]), "new_customers_per_month": A(50, [20, 60]), "fte_cost": A(0, 0)}},
    "project": {"revenue_unit": "project", "growth": G(motion="founder_led", ramp_k=1.0, ramp_mid_month=2, sales_cycle_months=1), "fte": 0,
                "assumptions": {"project_value": A(800, [450, 1100]), "projects_per_year": A(4, [1.5, 3]), "gross_margin": A(85, [85, 92]), "monthly_churn": A(4, [4, 7]),
                                "expansion": A(0, [0, 1]), "cac": A(300, [300, 700]), "new_customers_per_month": A(2, [1, 2]), "fte_cost": A(0, 0)}},
    "transactional": {"revenue_unit": "transactional", "growth": G(motion="partner", ramp_k=1.0), "fte": 0,
                      "assumptions": {"gmv_per_customer_month": A(4000, [2000, 5000]), "take_rate": A(12, [8, 15]), "gross_margin": A(80, 80), "monthly_churn": A(6, [5, 10]),
                                      "cac": A(250, [200, 400]), "new_customers_per_month": A(6, [3, 8]), "fte_cost": A(0, 0)}},
    "one_off": {"revenue_unit": "one_off", "growth": G(), "fte": 0, "repeat_interval_months": 3,
                "assumptions": {"pricing": A(100, 100), "gross_margin": A(80, 80), "repeat_rate": A(20, 20), "cac": A(120, 120), "new_customers_per_month": A(1, 1), "fte_cost": A(0, 0)}},
    "one_off_ok": {"revenue_unit": "one_off", "growth": G(), "fte": 0, "repeat_interval_months": 3,
                   "assumptions": {"pricing": A(59, 59), "gross_margin": A(95, 95), "repeat_rate": A(20, 20), "cac": A(30, 30), "new_customers_per_month": A(1, 1), "fte_cost": A(0, 0)}},
    "paid": {"revenue_unit": "subscription", "growth": G(motion="paid"), "fte": 0,
             "assumptions": {"pricing": A(40, 40), "gross_margin": A(80, 80), "monthly_churn": A(5, 5), "monthly_budget": A(3000, 3000), "new_customers_per_month": A(4, [4, 8]), "fte_cost": A(0, 0)}},
    "trial_upfront": {"revenue_unit": "subscription", "growth": G(motion="paid"), "fte": 0, "free_months": 3, "upfront_investment": 5000, "fixed_costs_monthly": 100,
                      "assumptions": {"pricing": A(10, 10), "gross_margin": A(80, 80), "monthly_churn": A(10, 10), "monthly_budget": A(200, 200), "new_customers_per_month": A(10, 10), "fte_cost": A(0, 0)}},
    "profitable_upfront": {"revenue_unit": "subscription", "growth": G(), "fte": 0, "upfront_investment": 5000,
                           "assumptions": {"pricing": A(1000, 1000), "gross_margin": A(80, 80), "monthly_churn": A(5, 5), "cac": A(10, 10), "new_customers_per_month": A(10, 10), "fte_cost": A(0, 0)}},
    "post_revenue": {"revenue_unit": "subscription", "growth": G(motion="plg"), "fte": 0, "starting_customers": 1640, "starting_mrr": 41000,
                     "assumptions": {"pricing": A(25, 25, "observed"), "gross_margin": A(80, 80), "monthly_churn": A(4, [3, 7], "observed"), "expansion": A(1, 1), "cac": A(50, [20, 100], "observed"),
                                     "new_customers_per_month": A(100, 100, "observed"), "fte_cost": A(0, 0)}},
    "oneoff_cap": {"revenue_unit": "one_off", "growth": G(), "fte": 0, "repeat_interval_months": 120,
                   "assumptions": {"pricing": A(59, 59), "gross_margin": A(95, 95), "repeat_rate": A(80, 80), "cac": A(100, 100), "new_customers_per_month": A(1, 1), "fte_cost": A(0, 0)}},
    "channels": {"revenue_unit": "subscription", "growth": G(), "fte": 0, "channels": [{"name": "a", "customers_per_month": 100, "cac": 10}, {"name": "b", "customers_per_month": 1, "cac": 1000}],
                 "assumptions": {"pricing": A(100, 100), "gross_margin": A(80, 80), "monthly_churn": A(20, 20), "expansion": A(25, 25), "cac": A(None, 40),
                                 "new_customers_per_month": A(101, 101), "fte_cost": A(0, 0)}},
    "filled": {"revenue_unit": "subscription", "growth": G(motion="plg", ramp_k=1.0), "fte": 0,
               "assumptions": {"pricing": A(None, [50, 80]), "gross_margin": A(80, 80), "monthly_churn": A(3, [3, 6]), "cac": A(150), "new_customers_per_month": A(50, [20, 60]), "fte_cost": A(0, 0)}},
}
for _k in list(CASES_IN): CASES_IN[_k] = dict(CASES_IN[_k], **FP)
CASES_IN["horizon12"] = dict(CASES_IN["subscription_growth"], months=12)
CASES_IN["churn99"] = {**CASES_IN["subscription"], "assumptions": {**CASES_IN["subscription"]["assumptions"], "monthly_churn": A(99, 99)}}


def with_assumption(case, **over):
    c = json.loads(json.dumps(case)); c["assumptions"].update(over); return c


def python_checks():
    # 1. cohort expansion: price 100, churn 20 %, expansion 25 %, 1 new/month → month-2 revenue 200 (q = 1)
    r = run_line(CASES_IN["subscription"])["cases"]["founder"]
    check("cohort expansion: month-2 revenue = 200 (not 225)", close(r["projection"][1]["mrr"], 200), r["projection"][1]["mrr"])
    check("q=1 LTV = contribution × 60", close(r["unit_economics"]["ltv"], 80 * 60), r["unit_economics"]["ltv"])
    check("q=1 payback = ceil(120/80) = 2", r["unit_economics"]["payback_months"] == 2, r["unit_economics"]["payback_months"])
    # 2. one-off
    r = run_line(CASES_IN["one_off"])["cases"]["founder"]
    check("one-off LTV ≈ 100·0.8/(1−0.2) with 60-month cap", close(r["unit_economics"]["ltv"], 100), r["unit_economics"]["ltv"])
    check("one-off LTV<CAC → payback None", r["unit_economics"]["payback_months"] is None)
    r = run_line(CASES_IN["one_off_ok"])["cases"]["founder"]
    check("one-off payback = 1 when first purchase covers CAC", r["unit_economics"]["payback_months"] == 1, r["unit_economics"]["payback_months"])
    r = run_line(CASES_IN["oneoff_cap"])["cases"]["founder"]
    check("one-off LTV with 120-month interval = single purchase contribution", close(r["unit_economics"]["ltv"], 59 * 0.95), r["unit_economics"]["ltv"])
    check("one-off: LTV < CAC ⇒ payback None (consistent with capped LTV)", r["unit_economics"]["ltv_cac"] < 1 and r["unit_economics"]["payback_months"] is None)
    # 3. two cases, bridge and market range
    res = run_line(CASES_IN["subscription_growth"])
    check("four cases present", set(res["cases"]) == set(CASES))
    F, M = res["cases"]["founder"], res["cases"]["market"]
    check("founder case uses founder churn 3 %", close(F["inputs"]["monthly_churn_pct"], 3)); check("market case uses market mid churn 4.5 %", close(M["inputs"]["monthly_churn_pct"], 4.5))
    check("pessimistic case uses worst edges (churn 6, cac 300, price 50)", close(res["cases"]["pessimistic"]["inputs"]["monthly_churn_pct"], 6) and close(res["cases"]["pessimistic"]["inputs"]["blended_cac"], 300) and close(res["cases"]["pessimistic"]["inputs"]["arpu0"], 50))
    check("optimistic case uses best edges (churn 3, cac 150, price 80)", close(res["cases"]["optimistic"]["inputs"]["monthly_churn_pct"], 3) and close(res["cases"]["optimistic"]["inputs"]["blended_cac"], 150) and close(res["cases"]["optimistic"]["inputs"]["arpu0"], 80))
    br = {b["domain"]: b for b in res["bridge"]}
    check("bridge has a row per differing domain, none for equal ones", set(br) == {"pricing", "monthly_churn", "expansion", "cac", "new_customers_per_month"}, sorted(br))
    check("bridge rows share ONE base: the market case", all(b["peak_capital_need"]["market"] == M["peak_capital_need"] for b in res["bridge"]))
    check("bridge ranked by |Δ capital| descending", [abs(b["peak_capital_need"]["delta"]) for b in res["bridge"]] == sorted([abs(b["peak_capital_need"]["delta"]) for b in res["bridge"]], reverse=True))
    # bridge on the market case with founder churn must equal a direct run with that value
    direct = run_line(with_assumption(CASES_IN["subscription_growth"], pricing=A(65, 65), monthly_churn=A(3, 3), expansion=A(1.25, 1.25), cac=A(225, 225), new_customers_per_month=A(40, 40)))["cases"]["founder"]
    check("bridge(churn) == direct run at market values with founder churn", close(br["monthly_churn"]["peak_capital_need"]["if_founder_right"], direct["peak_capital_need"]), (br["monthly_churn"]["peak_capital_need"]["if_founder_right"], direct["peak_capital_need"]))
    cmp_ = {c["domain"]: c for c in res["comparison"]}
    check("comparison: gap ratio for churn 4.5/3 = 1.5", close(cmp_["monthly_churn"]["gap_ratio"], 1.5), cmp_["monthly_churn"])
    check("comparison: equal values → no gap, no bridge row", cmp_["gross_margin"]["gap_ratio"] is None and "gross_margin" not in br)
    # 4. fills
    r = run_line(with_assumption(CASES_IN["subscription_growth"], pricing=A(None, [50, 80])))
    row = next(c for c in r["comparison"] if c["domain"] == "pricing")
    check("missing founder value → filled from market, flagged, no bridge row", row["fill"] == "filled_from_market" and row["founder"] == 65 and "pricing" not in {b["domain"] for b in r["bridge"]}, row)
    r = run_line(with_assumption(CASES_IN["subscription_growth"], pricing=A(70)))
    row = next(c for c in r["comparison"] if c["domain"] == "pricing")
    check("missing market value → founder used, flagged unverified", row["fill"] == "unverified" and row["market"] == 70, row)
    # 5. CAC: channels (founder) vs market cac; paid budget
    r = run_line(CASES_IN["channels"])
    check("channels: founder CAC = weighted 19.80", close(r["cases"]["founder"]["inputs"]["blended_cac"], 19.8), r["cases"]["founder"]["inputs"]["blended_cac"])
    check("channels: market CAC = market assumption 40 (bridge reaches the CAC actually used)", close(r["cases"]["market"]["inputs"]["blended_cac"], 40) and any(b["domain"] == "cac" for b in r["bridge"]), r["cases"]["market"]["inputs"]["blended_cac"])
    check("channels: comparison founder CAC == the CAC the founder case computed with", close(next(c for c in r["comparison"] if c["domain"] == "cac")["founder"], r["cases"]["founder"]["inputs"]["blended_cac"]))
    from financial_model import consistency
    for nm in ("subscription_growth", "project", "transactional", "one_off_ok", "paid", "channels", "post_revenue", "filled"):
        check(f"consistency: comparison ↔ cases ↔ bridge agree ({nm})", consistency(run_line(CASES_IN[nm])) == [], consistency(run_line(CASES_IN[nm])))
    r = run_line(CASES_IN["paid"])
    check("paid CAC = 3000/4 = 750 (founder); market 3000/6 = 500", close(r["cases"]["founder"]["inputs"]["blended_cac"], 750) and close(r["cases"]["market"]["inputs"]["blended_cac"], 500))
    check("paid: bridge on new customers changes derived CAC", any(b["domain"] == "new_customers_per_month" and b["ltv_cac"]["if_founder_right"] != b["ltv_cac"]["market"] for b in r["bridge"]))
    rp = run_line(with_assumption(CASES_IN["paid"], monthly_budget=A(3000, [3000, 5000])))
    check("paid: joint budget+customers bridge row present", any(b["kind"] == "joint" for b in rp["bridge"]), [b["domain"] for b in rp["bridge"]])
    # 6. determinism
    check("deterministic", json.dumps(run_all(CASES_IN["project"]), sort_keys=True) == json.dumps(run_all(CASES_IN["project"]), sort_keys=True))
    # 7. validation
    bads = [
        (with_assumption(CASES_IN["subscription"], pricing=A(-100, -100)), "negative price"),
        (with_assumption(CASES_IN["subscription"], cac=A("500", 500)), "string CAC"),
        (with_assumption(CASES_IN["subscription"], pricing=A(70, [50, 80], source="see churn.saas.nope.2026")), "source citing a non-existent library entry"),
        (with_assumption(CASES_IN["subscription"], pricing=A(0, 0)), "price 0"),
        (with_assumption(CASES_IN["subscription"], pricing=A(70, [80, 50])), "reversed market range"),
        (with_assumption(CASES_IN["subscription"], pricing={"founder": 70}), "founder value without status"),
        (dict(CASES_IN["subscription"], fte=3), "fte 3 with fte_cost 0"),
        (with_assumption(CASES_IN["subscription"], fte_cost=A(8000, 8000)), "fte_cost with fte 0"),
        (dict(CASES_IN["paid"], channels=[{"name": "a", "customers_per_month": 4, "cac": 1}]), "channel table under a paid motion"),
        (dict(CASES_IN["subscription"], growth=G(sales_cycle_months=12), months=12), "lag ≥ horizon"),
        (dict(CASES_IN["subscription"], growth=G(ramp_k=10, ramp_mid_month=100), months=120), "ramp that overflows a spreadsheet exp()"),
        (dict(CASES_IN["subscription"], growth=G(motion="magic")), "unknown motion"),
        (dict(CASES_IN["subscription_growth"], comparable_benchmark={"segment_size": 5000, "comparable_penetration_pct": [0, 5]}), "penetration low = 0"),
        (dict(CASES_IN["channels"], assumptions={**CASES_IN["channels"]["assumptions"], "cac": A(150, 40)}), "founder cac contradicting the channel table"),
        (dict(CASES_IN["subscription"], channels=[{"name": "a", "customers_per_month": 1}]), "channel without cac"),
        (with_assumption(CASES_IN["subscription"], monthly_churn=A(100, 100)), "churn 100 %"),
        (with_assumption(CASES_IN["subscription"], gross_margin=A(0, 0)), "margin 0 %"),
        (with_assumption(CASES_IN["subscription"], expansion=A(50, 50)), "expansion 50 %"),
        (with_assumption(CASES_IN["subscription"], new_customers_per_month=A(0, 0)), "zero capacity"),
        (with_assumption(CASES_IN["subscription"], pricing=A(100, [50, 80], source="")), "market without source"),
        (with_assumption(CASES_IN["subscription"], pricing=A(100, 100, status="confident")), "unknown status"),
        (with_assumption(CASES_IN["subscription"], bogus=A(1, 1)), "unknown domain"),
        (with_assumption(CASES_IN["paid"], cac=A(100, 100)), "cac on a paid motion"),
        (with_assumption(CASES_IN["one_off"], repeat_rate=A(100, 100)), "repeat 100 %"),
        (dict(CASES_IN["subscription"], growth=G(ramp_k=-1)), "negative ramp_k"),
        (dict(CASES_IN["subscription"], growth=G(sales_cycle_months=-3)), "negative lag"),
        (dict(CASES_IN["subscription"], growth=G(funnel=[{"step": "x"}])), "funnel step without source"),
        (dict(CASES_IN["subscription"], channels=[{"name": "a", "customers_per_month": 0, "cac": 50}]), "zero-volume channels"),
        (dict(CASES_IN["subscription"], channels=[{"name": "a", "customers_per_month": "10", "cac": 50}]), "string channel volume"),
        (dict(CASES_IN["subscription_growth"], comparable_benchmark={"segment_size": 0, "comparable_penetration_pct": [1, 3]}), "segment_size 0"),
        (dict(CASES_IN["subscription"], starting_customers=5), "starting_customers without starting_mrr"),
        (dict(CASES_IN["one_off"], starting_customers=5, starting_mrr=100), "starting base on one_off"),
        (dict(CASES_IN["trial_upfront"], free_months=61), "free_months 61"),
    ] + [(dict(CASES_IN["subscription"], months=m), f"months={m}") for m in (0, -1, 11, 121, 36.5)]
    for bad, label in bads:
        try:
            run_all(bad); check(f"refuses {label}", False, "no error raised")
        except ValueError:
            check(f"refuses {label}", True)
        except Exception as e:
            check(f"refuses {label} with ValueError (got {type(e).__name__})", False, str(e)[:120])
    # 8. free months + upfront
    r = run_line(CASES_IN["trial_upfront"])["cases"]["founder"]
    check("free_months: no revenue in months 1-3", all(x["mrr"] == 0 for x in r["projection"][:3]))
    check("free_months: month-4 revenue = 10 × 0.9^3 × 10 = 72.9", close(r["projection"][3]["mrr"], 72.9), r["projection"][3]["mrr"])
    check("upfront in peak capital (≥ 5000)", r["peak_capital_need"] >= 5000)
    # v24: sustained break-even, micro business, tornado, loaded CAC, cash-out, library flags
    live_dip = {"revenue_unit": "subscription", "growth": G(motion="paid", ramp_mid_month=6), "fte": 1, "starting_customers": 200, "starting_mrr": 4000,
                "assumptions": {"pricing": A(20, 20, "observed"), "gross_margin": A(80, 80), "monthly_churn": A(6, 6, "observed"), "monthly_budget": A(9000, 9000), "new_customers_per_month": A(60, 60, "observed"), "fte_cost": A(3000, 3000)}}
    rl = run_line(live_dip)["cases"]["founder"]
    nets = [x["net_cash_flow"] for x in rl["projection"]]
    check("model: break-even is sustained (a positive month 1 followed by losses is not break-even)", (rl["breakeven_month"] is None) == (min(nets[-1:]) < 0) and (rl["breakeven_month"] is None or all(n_ >= 0 for n_ in nets[rl["breakeven_month"] - 1:])), (rl["breakeven_month"], nets[:3], nets[-2:]))
    micro = {"revenue_unit": "subscription", "growth": G(), "fte": 0, "fixed_costs_monthly": 37,
             "assumptions": {"pricing": A(3.3, 3.3), "gross_margin": A(80, 80), "monthly_churn": A(5, 5), "cac": A(7.7, 7.7), "new_customers_per_month": A(2, 2), "fte_cost": A(0, 0)}}
    ok_micro = True
    try: run_line(micro)
    except AssertionError as e: ok_micro = False; msg_micro = str(e)
    check("model: micro business (capital need below €1k) does not trip the breakdown reconciliation", ok_micro, None if ok_micro else msg_micro)
    rt = run_line(CASES_IN["subscription_growth"])
    check("model: tornado has one row per ranged domain, sorted by capital width", all(rt["tornado"][i]["width_capital"] >= rt["tornado"][i + 1]["width_capital"] for i in range(len(rt["tornado"]) - 1)) and all(t["domain"] in {c_["domain"] for c_ in rt["comparison"]} for t in rt["tornado"]))
    loaded = dict(CASES_IN["subscription"], fte=2, sales_share_of_payroll=0.5, cash_on_hand=1000, assumptions={**CASES_IN["subscription"]["assumptions"], "fte_cost": A(4000, 4000)})
    rlo = run_line(loaded)["cases"]["market"]; ue_ = rlo["unit_economics"]
    check("model: loaded CAC = cash CAC + selling payroll per new customer; headline ratio uses it", abs(ue_["cac_loaded"] - (rlo["inputs"]["blended_cac"] + 2 * 4000 * 0.5 / rlo["inputs"]["new_customers_per_month_full"])) < 0.01 and ue_["ltv_cac_headline"] == ue_["ltv_cac_loaded"] and ue_["ltv_cac_loaded"] < ue_["ltv_cac"], ue_)
    check("model: cash_on_hand gives a cash-out month when cumulative cash falls below it", rlo["cash_out_month"] is not None and rlo["projection"][rlo["cash_out_month"] - 1]["cumulative"] < -1000 and any(f_["id"] == "cash_out" for f_ in rlo["flags"]), (rlo["cash_out_month"], rlo["flags"]))
    unv = dict(CASES_IN["subscription"], assumptions={**CASES_IN["subscription"]["assumptions"], "gross_margin": {"founder": 60, "market": [40, 65], "status": "founder_stated", "source": "margin.services (library)"}})
    ru_ = run_line(unv)
    gm = next(c_ for c_ in ru_["comparison"] if c_["domain"] == "gross_margin")
    check("model: an unverified / stale library entry is flagged on the row and in warnings", "unverified" in gm["library_flags"] and any("margin.services" in w for w in ru_["warnings"]), (gm["library_flags"], ru_["warnings"][:2]))
    bad_ccy = dict(CASES_IN["subscription"], currency="EUR", assumptions={**CASES_IN["subscription"]["assumptions"], "pricing": {"founder": 50, "market": [40, 80], "status": "founder_stated", "source": "x", "currency": "USD"}})
    check("model: a market value in another currency is refused", _raises(lambda: run_line(bad_ccy), "currency"))
    check("model: fx rate given as a string is refused, not crashed", _raises(lambda: fm.capital_vs_round("pre_seed", 100000, "EUR", {"USD_per_EUR": "abc"}), "fx_rates"))
    cvr_ = fm.capital_vs_round("pre_seed", 100000, "EUR", {"USD_per_EUR": 1.08})
    check("model: round bounds are exposed in the model currency", cvr_["round_low_model_ccy"] == round(250000 / 1.08, 0) and cvr_["round_high_model_ccy"] == round(1000000 / 1.08, 0), cvr_)
    r = run_line(CASES_IN["profitable_upfront"])["cases"]["founder"]
    check("profitable month 1: peak capital need = upfront 5000", close(r["peak_capital_need"], 5000), r["peak_capital_need"])
    c = run_all({"funding_path": "pre_seed", "lines": [dict(CASES_IN["profitable_upfront"], name="A"), dict(CASES_IN["profitable_upfront"], name="B")]})["combined"]["founder"]
    check("combined, both lines profitable in month 1: peak capital need = 10 000", close(c["peak_capital_need"], 10000), c["peak_capital_need"])
    # 9. starting base
    r = run_line(CASES_IN["post_revenue"])["cases"]["founder"]
    check("starting base: customers month 1 = 1640·0.96 + 100", close(r["projection"][0]["customers"], 1640 * 0.96 + 100), r["projection"][0]["customers"])
    check("starting base: revenue month 1 = 41000·0.96·1.01 + 100·25", close(r["projection"][0]["mrr"], 41000 * 0.96 * 1.01 + 2500), r["projection"][0]["mrr"])
    # 9b. measured mode: observed statuses switch the report's base to the measured case; bridge_to_market from it
    pr = run_line(CASES_IN["post_revenue"])
    check("measured mode detected from observed statuses", pr["mode"] == "measured" and pr["counts"]["mode"] == "measured" and pr["counts"]["starting_base"])
    btm = {b["domain"]: b for b in pr["bridge_to_market"]}
    check("bridge_to_market: base is the measured case, rows for differing domains only", btm and all(b["peak_capital_need"]["measured"] == pr["cases"]["founder"]["peak_capital_need"] for b in btm.values()) and "pricing" not in btm, sorted(btm))
    direct = run_line(with_assumption(CASES_IN["post_revenue"], monthly_churn=A(5, [3, 7], "observed")))["cases"]["founder"]
    check("bridge_to_market(churn) == direct run with churn at the category mid", close(btm["monthly_churn"]["revenue_end"]["if_at_category"], direct["milestones"]["end"]["mrr"]), (btm["monthly_churn"]["revenue_end"], direct["milestones"]["end"]["mrr"]))
    check("pre-revenue case: mode pre_revenue and no bridge_to_market", run_line(CASES_IN["subscription_growth"])["mode"] == "pre_revenue" and run_line(CASES_IN["subscription_growth"])["bridge_to_market"] == [])
    cat = run_line(with_assumption(CASES_IN["subscription_growth"], pricing=A(None, [50, 80]), monthly_churn=A(None, [3, 6]), cac=A(None, [150, 300]), new_customers_per_month=A(None, [20, 60])))
    check("category estimate flagged when half or more rows have no founder number", cat["counts"]["category_estimate"] is True and run_line(CASES_IN["subscription_growth"])["counts"]["category_estimate"] is False)
    # 10. horizon & comparable
    r12 = run_line(CASES_IN["horizon12"])
    check("months=12: milestones end at month 12, no m36", r12["cases"]["market"]["milestones"]["end"]["month"] == 12 and "m36" not in r12["cases"]["market"]["milestones"])
    check("months=12: bridge revenue_end.month = 12", all(b["revenue_end"]["month"] == 12 for b in r12["bridge"]))
    check("months=12: comparable (3-year reference) skipped with a note", "note" in (r12["cases"]["market"]["comparable_benchmark"] or {}))
    r48 = run_line(dict(CASES_IN["subscription_growth"], months=48))["cases"]["market"]
    check("months=48: comparable compares month 36", close(r48["comparable_benchmark"]["bottom_up_customers_m36"], r48["projection"][35]["customers"]))
    # 11. multi-line
    L1 = dict(CASES_IN["subscription_growth"], name="A", upfront_investment=5000); L2 = dict(CASES_IN["one_off_ok"], name="B", upfront_investment=7000)
    MM = run_all({"funding_path": "bootstrap", "currency": "EUR", "lines": [L1, L2]}); c = MM["combined"]["market"]
    check("combined peak capital ≥ 12 000 (both upfronts)", c["peak_capital_need"] >= 12000, c["peak_capital_need"])
    check("combined is order-independent", run_all({"funding_path": "seed", "lines": [L2, L1]})["combined"]["market"]["peak_capital_need"] == c["peak_capital_need"])
    check("combined bridge present, base = combined market capital, ranked", MM["combined"]["bridge"] and all(x["peak_capital_need"]["market"] == c["peak_capital_need"] for x in MM["combined"]["bridge"])
          and [abs(x["peak_capital_need"]["delta"]) for x in MM["combined"]["bridge"]] == sorted([abs(x["peak_capital_need"]["delta"]) for x in MM["combined"]["bridge"]], reverse=True))
    check("capital_vs_round: verified bootstrap entry in EUR, position stated", MM["capital_vs_round"]["position"] in ("below_round", "inside_round", "above_round") and MM["capital_vs_round"]["round_high"] == 50000, MM["capital_vs_round"])
    us = run_all(CASES_IN["subscription_growth"])["capital_vs_round"]
    check("capital_vs_round: pre_seed in USD uses the verified Carta entry", us["entry"] == "round.pre_seed.us.2025" and us["position"] in ("below_round", "inside_round", "above_round"), us)
    eu = run_all(dict(CASES_IN["subscription_growth"], currency="EUR"))["capital_vs_round"]
    check("capital_vs_round: pre_seed in EUR without a rate → unavailable with reason", eu["position"] == "unavailable" and "fx_rates" in eu["reason"], eu)
    eu2 = run_all(dict(CASES_IN["subscription_growth"], currency="EUR", fx_rates={"USD_per_EUR": 1.1}))["capital_vs_round"]
    check("capital_vs_round: pre_seed in EUR with a rate → converted and positioned", eu2["position"] != "unavailable" and eu2["fx_rate"] == 1.1 and close(eu2["capital_in_round_currency"], eu2["capital_need"] * 1.1), eu2)
    sd = run_all(dict(CASES_IN["subscription_growth"], funding_path="seed"))["capital_vs_round"]
    check("capital_vs_round: series_a has no verified entry → unavailable", run_all(dict(CASES_IN["subscription_growth"], funding_path="series_a"))["capital_vs_round"]["position"] == "unavailable")
    check("capital_vs_round: seed — a verified EUR entry exists but the model is in USD → unavailable with the currency reason", sd["position"] == "unavailable" and "EUR" in sd["reason"], sd)
    nc = run_all({k: v for k, v in CASES_IN["subscription_growth"].items() if k != "currency"})["capital_vs_round"]
    check("capital_vs_round: missing currency → unavailable", nc["position"] == "unavailable" and "currency" in nc["reason"], nc)
    check("counts aggregated over lines", MM["counts"]["total"] == sum(ln["counts"]["total"] for ln in MM["lines"].values()))
    for bad, label in [({"funding_path": "seed", "currency": "EUR", "lines": [dict(L1, months=12), dict(L2, months=36)]}, "mismatched horizons"), ({"funding_path": "seed", "currency": "EUR", "lines": [dict(L1, name="Basic"), dict(L2, name="basic")]}, "line names differing only by case"),
                       ({"lines": [L1, L2], "currency": "EUR"}, "missing funding_path"), (dict(CASES_IN["subscription"], funding_path="angel"), "unknown funding_path")]:
        try:
            run_all(bad); check(f"refuses {label}", False, "no error")
        except ValueError:
            check(f"refuses {label}", True)
    # 12. churn 99 %
    r = run_line(CASES_IN["churn99"])["cases"]["founder"]
    check("churn 99 % computes (LTV ≈ contribution)", close(r["unit_economics"]["ltv"], 80 * (1 + 0.0125 + 0.0125 ** 2), 0.05), r["unit_economics"]["ltv"])
    # payback at an exact boundary: CAC = 7 + 7·0.96 → 2 months, not 3
    rb = run_line(with_assumption(CASES_IN["subscription"], pricing=A(10, 10), gross_margin=A(70, 70), monthly_churn=A(4, 4), expansion=A(0, 0), cac=A(13.72, 13.72)))["cases"]["founder"]
    check("payback at exact float boundary = 2", rb["unit_economics"]["payback_months"] == 2, rb["unit_economics"]["payback_months"])
    single = run_all(CASES_IN["subscription_growth"])
    check("single-line: capital_vs_round and counts present", single["capital_vs_round"]["path"] == "pre_seed" and single["counts"]["total"] == 7, (single["capital_vs_round"], single["counts"]))
    # 1. sources with URLs are fine; a bogus library id is not; a domain is not mistaken for an id
    ok_url = with_assumption(CASES_IN["subscription_growth"], pricing=A(70, [50, 80], source="https://www.stripe.com/pricing and price.validation_tools.2026"))
    check("source with a URL and a real library id accepted", run_line(ok_url)["comparison"][0]["source"].startswith("https"))
    # 5. cash CAC 0: model runs, LTV/CAC n/a, payback 1; founder time cost gives the incl-time ratios without touching cash
    z = run_line(with_assumption(dict(CASES_IN["subscription"], founder_time_cost_per_customer=500), cac=A(0, 0)))["cases"]["founder"]
    check("cash CAC 0: ltv_cac None, payback 1, no acquisition cash flow", z["unit_economics"]["ltv_cac"] is None and z["unit_economics"]["payback_months"] == 1 and z["unit_economics"]["cash_cac_zero"], z["unit_economics"])
    check("cash CAC 0: capital need = 0 (founder time is not cash)", z["peak_capital_need"] == 0, z["peak_capital_need"])
    check("founder time cost → ltv_cac_incl_time = LTV/500 and payback_incl_time > 1", close(z["unit_economics"]["ltv_cac_incl_time"], z["unit_economics"]["ltv"] / 500) and (z["unit_economics"]["payback_incl_time"] or 99) > 1, z["unit_economics"])
    # 2. CLI without --json must not crash on a joint row
    import subprocess, tempfile as _tf
    with _tf.TemporaryDirectory() as _d:
        pj = os.path.join(_d, "p.json"); json.dump(with_assumption(CASES_IN["paid"], monthly_budget=A(3000, [3000, 5000])), open(pj, "w"))
        for flag in ([], ["--json"]):
            out = subprocess.run([sys.executable, os.path.join(HERE, "financial_model.py"), pj] + flag, capture_output=True, text=True)
            check(f"CLI {' '.join(flag) or 'markdown'} runs on a model with a joint bridge row", out.returncode == 0 and ("joint" in out.stdout or "together" in out.stdout), out.stderr[-200:])
    report_checks()


# ── report builder ───────────────────────────────────────────

def _fake_charts(d, model, br_rows=True, line=None, lang="en", cur="$"):
    import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
    from build_workbook import write_manifest
    files = ["chart_mrr.png", "chart_cash.png", "chart_customers.png", "dashboard.png"] + (["bridge.png"] if br_rows else [])
    for f in files:
        fig = plt.figure(figsize=(2, 1)); fig.savefig(os.path.join(d, f)); plt.close(fig)
    write_manifest(d, "charts.json", model, files[:3], {"line": None, "lang": lang, "currency": cur}); write_manifest(d, "dashboard.json", model, files[3:], {"line": line if "lines" in model else None, "lang": lang, "currency": cur})


LONG = "This sentence is long enough to pass the minimum-content check of the builder."
PARA = LONG + " " + LONG


def stakes_for(d):
    return "Capital need {{cases.market.peak_capital_need|money}} → {{bridge_by_domain.%s.peak_capital_need.if_founder_right|money}} if the team is right about this number." % d


def _minimal_report(**over):
    sec = lambda: {"lead": LONG, "paragraphs": [PARA]}
    stakes = "Capital need {{cases.market.peak_capital_need|money}} → {{bridge_by_domain.%s.peak_capital_need.if_founder_right|money}} if the team is right about this number."
    rep = {"startup": "Acme", "date": "2026-01-01", "lang": "en", "currency": "$", "verdict": "Proceed with conditions", "conditions": ["3 paying customers within 90 days", "churn under 5 % for 3 months"],
           "verdict_text": "Verdict text: {{capital_vs_round.position|text}} against the stated funding path; the market case decides.",
           "story": [LONG, LONG, LONG, LONG, LONG], "evidence_basis": "Of {{counts.total|int}} assumptions: {{counts.observed|int}} measured.",
           "assumptions": {"lead": LONG, "stories": [
               {"domain": d, "belief": LONG, "evidence": LONG, "stakes": stakes % d, "test": "Measure it over 30 days; pass if above the market mid, fail otherwise."} for d in ("monthly_churn", "cac", "new_customers_per_month", "pricing", "expansion")]},
           "market_case": sec(), "founder_case": sec(),
           "competition": {"lead": LONG, "paragraphs": [PARA], "where_we_win": LONG,
                           "dimensions": [{"key": "a", "label": "Feature A", "why": LONG}, {"key": "b", "label": "Feature B", "why": LONG}],
                           "rows": [{"name": "Self", "layer": "direct", "price": "n/a", "cells": {"a": "yes", "b": "no"}, "source": "intake", "self": True},
                                    {"name": "X", "layer": "direct", "price": "n/a", "cells": {"a": "yes", "b": "yes"}, "source": "site"},
                                    {"name": "Y", "layer": "indirect", "price": "n/a", "cells": {"a": "no", "b": "yes"}, "source": "site"},
                                    {"name": "Z", "layer": "displacer", "price": "n/a", "cells": {"a": "no", "b": "no"}, "source": "site"},
                                    {"name": "W", "layer": "direct", "price": "n/a", "cells": {"a": "partly", "b": "no"}, "source": "site"}]},
           "next_90_days": [{"title": "a", "what": LONG, "how": LONG, "threshold": "At least 20 paying customers within 60 days.", "domain": d} for d in ("monthly_churn", "cac", "new_customers_per_month")],
           "risks": [{"title": "risk one", "what": LONG, "impact": "{{months|int}} " + LONG, "signal": "3 months of the metric below the line", "action": LONG, "likelihood": "High", "reversibility": "Reversible"},
                     {"title": "risk two", "what": LONG, "impact": "{{counts.total|int}} " + LONG, "signal": "2 lost accounts in 1 quarter", "action": LONG, "likelihood": "Low", "reversibility": "Structural"}],
           "appendix": {"evidence": [["#", "Claim"], ["one", "x"]]},
           "manual_numbers": [{"value": "30", "field": "test window, days", "source": "proposed test", "where": ["assumptions.stories"], "context": "over 30 days"},
                              {"value": "20", "field": "threshold customers", "source": "proposed test", "where": ["next_90_days"], "context": "20 paying customers"},
                              {"value": "60", "field": "threshold days", "source": "proposed test", "where": ["next_90_days"], "context": "within 60 days"},
                              {"value": "90", "field": "condition window, days", "source": "proposed test", "where": ["conditions"], "context": "within 90 days"},
                              {"value": "5", "field": "condition churn, %", "source": "proposed test", "where": ["conditions"], "context": "under 5 %"}]}
    rep.update(over); return rep


def _build(br, rep, d, out, model):
    """Run br.build; a SystemExit becomes a failed check with the builder's message instead of killing the suite."""
    argv = sys.argv; sys.argv = ["x"]; err = io.StringIO()
    try:
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(err): br.build(rep, d, out, model)
        return True, ""
    except SystemExit:
        return False, err.getvalue()[-600:]
    finally:
        sys.argv = argv


def report_checks():
    import importlib
    br = importlib.import_module("build_report")
    model = run_all(CASES_IN["subscription_growth"])
    with tempfile.TemporaryDirectory() as d:
        _fake_charts(d, model)
        T = br.LABELS["en"]
        good = _minimal_report()
        check("report: placeholders in minimal report resolve", br.check_placeholders(good, model) == [], br.check_placeholders(good, model))
        check("report: free text without raw numbers passes", br.verify_free_text(good) == [], br.verify_free_text(good))
        rendered = br.render_placeholders(good, model, "$", T)
        check("report: minimal valid report passes requirements", br.check_requirements(rendered, model, d, T) == [], br.check_requirements(rendered, model, d, T))
        # placeholders
        ph = _minimal_report(market_case={"lead": "LTV/CAC is {{cases.founder.unit_economics.ltv_cac|mult}} vs {{cases.market.unit_economics.ltv_cac|mult}}.", "paragraphs": ["T."]})
        check("report: placeholder numbers pass verification", br.verify_free_text(ph) == [])
        rv = br.render_placeholders(ph, model, "$", T)["market_case"]["lead"]
        from chart_style import fmt_mult
        check("report: placeholder renders founder and market values", rv == f"LTV/CAC is {fmt_mult(model['cases']['founder']['unit_economics']['ltv_cac'])} vs {fmt_mult(model['cases']['market']['unit_economics']['ltv_cac'])}.", rv)
        for txt, what in [("{{cases.market.unit_economics.nope|mult}}", "missing field"), ("{{cases.market.unit_economics.ltv_cac|bad}}", "unknown format"),
                          ("value {{ here", "stray braces"), ("{{cases.market.unit_economics.ltv_cac|pct}}", "format mismatch (ratio as %)"),
                          ("{{cases.market.inputs.monthly_churn_pct|money}}", "format mismatch (% as money)"), ("{{cases.market.projection|text}}", "ref to a structure"),
                          ("{{cases.market.peak_capital_need|text}}", "money field as bare text"), ("{{cases.market.growth.motion|mult}}", "string field as a multiple"),
                          ("{{cases.market.inputs.monthly_churn_pct|text}}", "pct field as bare text")]:
            bad = _minimal_report(market_case={"lead": txt + " " + LONG, "paragraphs": [PARA]})
            check(f"report: placeholder {what} rejected", br.check_placeholders(bad, model) != [], txt)
        # placeholders are checked in every key family
        for where, bad in [("kpis.label", _minimal_report(kpis=[{"label": "LTV {{cases.market.unit_economics.nope|mult}}", "ref": "cases.market.unit_economics.ltv_cac", "fmt": "mult"}] * 3)),
                           ("appendix", _minimal_report(appendix={"evidence": [["#", "C"], ["one", "{{cases.market.x|bad}}"]]})),
                           ("comparison.notes", _minimal_report(comparison={"notes": {"pricing": "{{ oops"}})),
                           ("labels", _minimal_report(labels={"dashboard_note": "{{combined.market.peak_capital_need|money}} — {{nope}}"})),
                           ("startup", _minimal_report(startup="Acme {{cases.market.peak_capital_need|money}}"))]:
            check(f"report: bad placeholder in {where} rejected", br.check_placeholders(bad, model) != [])
        # numbers in prose
        for txt, expect in [("LTV/CAC is 999×.", 1), ("LTV/CAC is 999x, payback 480mo, TAM 12bn, 1e9 users.", 4), ("Costs $1999.", 1), ("Reach 2026× of baseline.", 1),
                            ("In 2026, three tests.", 0), ("Price $10, 300 seats.", 2), ("Takes 30 days.", 1), ("Three of 5 tests.", 0),
                            ("Costs 1.000 euro and 9.999 more.", 2), ("About 5 тыс. евро.", 1), ("Revenue in M36 and Q4.", 2), ("Ratio x45.", 1)]:
            bad = _minimal_report(market_case={"lead": txt + " " + LONG, "paragraphs": [PARA]})
            n = len(br.verify_free_text(bad)); check(f"report: prose {txt!r} → {expect} undeclared", n == expect, br.verify_free_text(bad))
        raw = _minimal_report(market_case={"lead": f"LTV/CAC is {model['cases']['market']['unit_economics']['ltv_cac']}×. " + LONG, "paragraphs": [PARA]})
        check("report: a raw model number typed into text is rejected", len(br.verify_free_text(raw)) == 1)
        for where, bad in [("kpis.label", _minimal_report(kpis=[{"label": "LTV/CAC 999× (mid)", "ref": "cases.market.unit_economics.ltv_cac", "fmt": "mult"}] * 3)),
                           ("evidence_basis", _minimal_report(evidence_basis="Of 14 assumptions: 0 observed.")),
                           ("appendix", _minimal_report(appendix={"evidence": [["#", "C"], ["cac", "$999k"]]})),
                           ("comparison.notes", _minimal_report(comparison={"notes": {"pricing": "benchmark 45 %"}})),
                           ("labels", _minimal_report(labels={"comparison": "Numbers (LTV/CAC 999×, capital $5M)"})),
                           ("startup", _minimal_report(startup="Acme 2.0 ($5M ARR)")),
                           ("risks[].notes", _minimal_report(risks=[{"title": "risk", "text": LONG, "likelihood": "High", "reversibility": "Reversible", "notes": "CAC is $999k"}] * 2))]:
            check(f"report: undeclared number in {where} rejected", br.verify_free_text(bad) != [])
        decl = _minimal_report(market_case={"lead": "LTV/CAC is 999×. " + LONG, "paragraphs": [PARA]}, manual_numbers=good["manual_numbers"] + [{"value": "999", "field": "deck claim", "source": "deck p.7", "where": ["market_case.lead"], "context": "LTV/CAC is 999×"}])
        check("report: declared manual number passes where and in the phrase it is declared", br.verify_free_text(decl) == [], br.verify_free_text(decl))
        elsewhere = _minimal_report(market_case={"lead": LONG, "paragraphs": ["LTV/CAC is 999×. " + PARA]}, manual_numbers=decl["manual_numbers"])
        check("report: the same number outside its declared 'where' is rejected", br.verify_free_text(elsewhere) != [])
        other_meaning = _minimal_report(market_case={"lead": "Capital need is $999. " + LONG, "paragraphs": [PARA]}, manual_numbers=decl["manual_numbers"])
        check("report: the same value with another meaning (different phrase) is rejected", br.verify_free_text(other_meaning) != [])
        for bad_mn, what in [(good["manual_numbers"] + [{"value": "999", "field": "f", "source": "s", "context": "999×"}], "without where"),
                             (good["manual_numbers"] + [{"value": "999", "field": "f", "source": "s", "where": ["market_case"]}], "without context"),
                             (good["manual_numbers"] + [{"value": "999", "field": "f", "source": "s", "where": ["market_case"], "context": "no number here"}], "context lacking the number"),
                             (good["manual_numbers"] + [{"value": "999", "field": "f", "source": "model bridge.0", "where": ["market_case"], "context": "999×"}], "sourced from the model"), (["999"], "not an object")]:
            try:
                br.verify_free_text(dict(decl, manual_numbers=bad_mn)); check(f"report: manual number {what} rejected", False, "no error")
            except ValueError:
                check(f"report: manual number {what} rejected", True)
        # requirements
        R = lambda **o: br.check_requirements(br.render_placeholders(_minimal_report(**o), model, "$", T), model, d, T)
        check("report: verdict outside the three rejected", any("verdict" in e for e in R(verdict="Launch immediately")))
        check("report: 0 experiments rejected", any("next_90_days" in e for e in R(next_90_days=[])))
        check("report: 6 experiments rejected", any("next_90_days" in e for e in R(next_90_days=good["next_90_days"] * 2)))
        check("report: whitespace-only lead rejected", any("market_case.lead" in e for e in R(market_case={"lead": "   ", "paragraphs": [PARA]})))
        check("report: one-word paragraph rejected", any("paragraphs[0]" in e for e in R(market_case={"lead": LONG, "paragraphs": ["."]})))
        check("report: competition without a displacer row rejected", any("displacer" in e for e in R(competition={**good["competition"], "rows": [r for r in good["competition"]["rows"] if r["layer"] != "displacer"] + [dict(good["competition"]["rows"][1], name="X2")]})))
        check("report: competition with two self rows rejected", any("self" in e for e in R(competition={**good["competition"], "rows": [dict(r, self=True) for r in good["competition"]["rows"]]})))
        check("report: competition row missing a dimension cell rejected", any("no cell for dimension" in e for e in R(competition={**good["competition"], "rows": [dict(good["competition"]["rows"][0], cells={"a": "yes"})] + good["competition"]["rows"][1:]})))
        check("report: one competition dimension rejected", any("dimensions" in e for e in R(competition={**good["competition"], "dimensions": good["competition"]["dimensions"][:1], "rows": [dict(r, cells={"a": r["cells"]["a"]}) for r in good["competition"]["rows"]]})))
        check("report: risk impact without a number or placeholder rejected", any("impact" in e for e in R(risks=[dict(good["risks"][0], impact=LONG)] + good["risks"][1:])))
        check("report: risk signal without a number rejected", any("signal" in e for e in R(risks=[dict(good["risks"][0], signal="something goes wrong at some point")] + good["risks"][1:])))
        check("report: story with four sentences rejected", any("story" in e for e in R(story=[LONG] * 4)))
        check("report: story of dots rejected", any("story" in e for e in R(story=[".", ".", ".", ".", "."])))
        st_ok = good["assumptions"]["stories"]
        check("report: story missing the top bridge domain rejected", any("stories must cover" in e for e in R(assumptions={"lead": LONG, "stories": [s for s in st_ok if s["domain"] == "expansion"]})))
        check("report: story part too short rejected", any("stakes must" in e or "stakes must be" in e for e in R(assumptions={"lead": LONG, "stories": [dict(s, stakes="S.") for s in st_ok]})))
        check("report: story for a domain not in the model rejected", any("not in the model" in e for e in R(assumptions={"lead": LONG, "stories": st_ok + [dict(st_ok[0], domain="unicorns")]})))
        check("report: next_90_days without a test for the top bridge domain rejected", any("largest effect" in e for e in R(next_90_days=[dict(x, domain="other") for x in good["next_90_days"]])))
        check("report: threshold without a number rejected", any("threshold must state" in e for e in R(next_90_days=[dict(x, threshold="Enough customers within a while.") for x in good["next_90_days"]])))
        check("report: risk without likelihood enum rejected", any("likelihood" in e for e in R(risks=[dict(r, likelihood="maybe") for r in good["risks"]])))
        check("report: malformed report (string instead of list) fails with a message", br.check_schema(_minimal_report(competition={**good["competition"], "rows": "Acme Corp"})) != [])
        check("report: unequal evidence rows fail with a message", br.check_schema(_minimal_report(appendix={"evidence": [["a", "b"], ["c"]]})) != [])
        okv, msgv = _build(br, _minimal_report(verdict_text="Verdict without the funding-path line, long enough to pass the length check."), d, os.path.join(d, "v.docx"), model)
        check("report: verdict without the funding-path placeholder is refused at build", not okv and "capital_vs_round" in msgv, msgv[-200:])
        wrong = [dict(s, stakes="Capital need {{cases.market.peak_capital_need|money}} → {{bridge.0.peak_capital_need.if_founder_right|money}} if right about this.") for s in good["assumptions"]["stories"]]
        errs_w = br.check_story_bindings(_minimal_report(assumptions={"lead": LONG, "stories": wrong}), model)
        check("report: a story quoting another assumption's bridge row (bridge.0 for every domain) is rejected", sum("effect of another assumption" in e for e in errs_w) >= 3, errs_w[:3])
        wrong2 = [dict(s, stakes="Capital {{bridge_by_domain.cac.peak_capital_need.if_founder_right|money}} if the team is right.") if s["domain"] == "monthly_churn" else s for s in good["assumptions"]["stories"]]
        check("report: churn story quoting the CAC effect by name is rejected", any("effect of another assumption" in e for e in br.check_story_bindings(_minimal_report(assumptions={"lead": LONG, "stories": wrong2}), model)))
        check("report: stories bound to their own domain pass", br.check_story_bindings(good, model) == [], br.check_story_bindings(good, model))
        okw, msgw = _build(br, _minimal_report(assumptions={"lead": LONG, "stories": wrong}), d, os.path.join(d, "w.docx"), model)
        check("report: build refuses the mis-attributed stories", not okw and "another assumption" in msgw, msgw[-200:])
        check("report: verdict_text over 350 chars rejected", any("page-1 budget" in e for e in R(verdict_text="x" * 351)))
        check("report: KPI ref to a missing field rejected", any("kpi" in e for e in R(kpis=[{"label": "L", "ref": "cases.market.unit_economics.typo", "fmt": "mult"}] * 3)))
        check("report: KPI fmt mismatching the field rejected", any("does not fit" in e for e in R(kpis=[{"label": "L", "ref": "cases.market.unit_economics.ltv_cac", "fmt": "pct"}] * 3)))
        check("report: KPI with bogus fmt rejected", any("unknown fmt" in e for e in R(kpis=[{"label": "L", "ref": "cases.market.unit_economics.ltv_cac", "fmt": "bogus"}] * 3)))
        check("report: invalid labels override rejected", br.check_labels(_minimal_report(labels={"verdict_names": {"Proceed": "GO"}}), T) != [] and br.check_labels(_minimal_report(labels={"dashboard": ""}), T) != [] and br.check_labels(_minimal_report(labels={"nope": "x"}), T) != [])
        nofunnel = run_all(dict(CASES_IN["subscription_growth"], growth=G(motion="plg", ramp_k=1.0, funnel=[])))
        with tempfile.TemporaryDirectory() as d2:
            _fake_charts(d2, nofunnel); check("report: empty funnel rejected", any("funnel is empty" in e for e in br.check_requirements(br.render_placeholders(_minimal_report(), nofunnel, "$", T), nofunnel, d2, T)))
        same = run_all(with_assumption(CASES_IN["subscription"]))   # founder == market everywhere, nothing observed
        with tempfile.TemporaryDirectory() as d3:
            _fake_charts(d3, same, br_rows=False)
            nob = _minimal_report(assumptions={"lead": LONG, "stories": []})
            for s_ in nob["assumptions"]["stories"]: s_["stakes"] = LONG + " 1"
            check("report: no differing assumption and nothing measured → rejected", any("nothing to compare" in e for e in br.check_requirements(br.render_placeholders(nob, same, "$", T), same, d3, T)))
            check("report: a model without a bridge does not require bridge.png", "bridge.png" not in br.required_charts(same, {})["dashboard.json"])
        live = run_all(CASES_IN["post_revenue"])
        with tempfile.TemporaryDirectory() as d4:
            _fake_charts(d4, live)
            stakes_m = lambda d: "Revenue at the last month {{cases.founder.milestones.end.mrr|money}} → {{bridge_to_market_by_domain.%s.revenue_end.if_at_category|money}} if this metric reached the category level." % d
            st_l = [dict(good["assumptions"]["stories"][0], domain=d, stakes=stakes_m(d)) for _, d in br.top_bridge_domains(live, {})]
            nx_l = [dict(good["next_90_days"][0], domain=d) for _, d in br.top_bridge_domains(live, {}, k=1)] + good["next_90_days"][:2]
            rep_m = _minimal_report(assumptions={"lead": LONG, "stories": st_l}, next_90_days=nx_l)
            check("report: measured-mode stories bound to bridge_to_market pass the binding check", br.check_story_bindings(rep_m, live) == [], br.check_story_bindings(rep_m, live))
            wrong_m = _minimal_report(assumptions={"lead": LONG, "stories": [dict(s, stakes=stakes_m("cac")) for s in st_l]})
            check("report: measured-mode story quoting another metric's to-market effect rejected", any("another assumption" in e for e in br.check_story_bindings(wrong_m, live)))
            Tm = dict(T); Tm.update(T["_measured"])
            errs = br.check_requirements(br.render_placeholders(rep_m, live, "$", Tm), live, d4, Tm)
            check("report: post-revenue case with measured numbers passes requirements", errs == [], errs)
            okm, msgm = _build(br, rep_m, d4, os.path.join(d4, "m.docx"), live); check("report: measured-mode report builds", okm, msgm)
            from docx import Document as _Doc
            heads_m = [p.text for p in _Doc(os.path.join(d4, "m.docx")).paragraphs if p.style.name.startswith("Heading")]
            check("report: measured-mode headings ('Where you stand', not 'If the team is right')", T["_measured"]["comparison"] in heads_m and T["founder_case"] not in heads_m, heads_m[:6])
        # charts
        other = run_all(with_assumption(CASES_IN["subscription_growth"], pricing=A(71, [50, 80])))
        check("report: charts from a different model rejected", any("different model" in e for e in br.check_charts(d, other, good)))
        with tempfile.TemporaryDirectory() as d5:
            _fake_charts(d5, model)
            import matplotlib.pyplot as plt
            fig = plt.figure(figsize=(3, 1)); fig.savefig(os.path.join(d5, "dashboard.png")); plt.close(fig)
            check("report: replaced dashboard.png rejected", any("hash differs" in e for e in br.check_charts(d5, model, good)))
            st5 = json.load(open(os.path.join(d5, "dashboard.json"))); st5["context"]["line"] = "Other"; json.dump(st5, open(os.path.join(d5, "dashboard.json"), "w"))
            check("report: hand-edited manifest context rejected (signature)", any("signature" in e for e in br.check_charts(d5, model, good)))
            json.dump({"model_hash": st5["model_hash"], "files": [], "png_hash": {}}, open(os.path.join(d5, "dashboard.json"), "w"))
            check("report: manifest with an empty file list rejected (required set is fixed)", any("dashboard.png" in e or "signature" in e for e in br.check_charts(d5, model, good)))
        # lint
        hits = br.lint_register(_minimal_report(market_case={"lead": "The US market is the named mechanism.", "paragraphs": [PARA]}, next_90_days=[dict(x, title="Will users pay?") for x in good["next_90_days"]]))
        check("report: lint — 'US', 'named' and a question in an experiment title are not flagged", hits == [], hits)
        check("report: lint — second person still flagged", any(l == "second person" for _, l, _ in br.lint_register(_minimal_report(market_case={"lead": "You should test.", "paragraphs": [PARA]}))))
        check("report: lint — 'I think' / 'Мне кажется' / mid-paragraph question flagged", all(br.lint_register(_minimal_report(market_case={"lead": x, "paragraphs": [PARA]})) for x in ("I think this works.", "Мне кажется, это работает.", "Is it viable? Yes.")))
        # build + page order + pagination
        out = os.path.join(d, "t.docx")
        ok_b, msg_b = _build(br, good, d, out, model); check("report: minimal report builds", ok_b, msg_b)
        try:
            import weasyprint  # noqa: F401
            out_pdf = os.path.join(d, "r.pdf")
            ok_p, msg_p = _build(br, good, d, out_pdf, model); check("report: PDF renderer builds and passes the pagination check", ok_p and os.path.exists(out_pdf), msg_p)
            if ok_p:
                import subprocess as _sp
                txt1 = _sp.run(["pdftotext", "-f", "1", "-l", "1", out_pdf, "-"], capture_output=True, text=True).stdout
                check("report: PDF page 1 carries the verdict and the story", good["startup"] in txt1 and "risk one" not in txt1)
                txt = _sp.run(["pdftotext", out_pdf, "-"], capture_output=True, text=True).stdout
                check("report: PDF carries the competitive table and the risk table", "Feature A" in txt and "risk one" in txt and "Self" in txt)
        except ImportError:
            print("SKIP report: PDF renderer (weasyprint not installed)")
        out_html = os.path.join(d, "r.html")
        ok_h, msg_h = _build(br, good, d, out_html, model); check("report: HTML renderer builds", ok_h and os.path.exists(out_html), msg_h)
        if ok_h:
            ht = open(out_html, encoding="utf-8").read()
            check("report: HTML carries the financial picture (P&L, cash, capital, signals) as inline SVG", ht.count("<svg") >= 3 and 'class="t flags"' in ht)
            check("report: HTML is self-contained (fonts inline, no external URLs)", "data:font/ttf" in ht and "http://" not in ht.replace("http://www.w3.org", "") and "https://" not in ht)
        for c in ("founder", "market"):
            bd = model["cases"][c]["capital_breakdown"]
            recon = bd["upfront"] + bd["payroll"] + bd["marketing"] + bd["fixed"] - bd["gross_profit_offset"]
            check(f"model: capital breakdown reconciles with peak capital need ({c})", abs(recon - model["cases"][c]["peak_capital_need"]) <= 2 or model["cases"][c]["peak_capital_need"] == 0)
        from docx import Document
        doc = Document(out); heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        order = [heads.index(T["comparison"]), heads.index(T["dashboard"]), heads.index(T["assumptions"]), heads.index(T["market_case"]), heads.index(T["next_90"])]
        check("report: section order story/comparison → dashboard → assumptions → market case → next 90 days", order == sorted(order), heads[:8])
        ok, msg = br.verify_pagination(out, T); check("report: rendered pagination — comparison on page 1, dashboard on page 2", ok is True, msg)
        body = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        from chart_style import fmt_mult as _fm
        check("report: KPI tile shows founder | market LTV/CAC", f"{_fm(model['cases']['founder']['unit_economics']['ltv_cac'])} | {_fm(model['cases']['market']['unit_economics']['ltv_cac'])}" in body)
        check("report: funding-path position rendered as words, not an enum token", "above_round" not in body and "below_round" not in body and "inside_round" not in body)
        check("report: LTV row named as gross-margin contribution", T["ue_ltv"] in body)
        check("report: comparison table lists the market source", "test source" in body)
        check("report: assumption stories rendered with their four parts", body.count(T["st_stakes"]) >= 3)
        # multi-line: line required; many rows → page 1 keeps 6, pagination holds
        mm = run_all({"funding_path": "bootstrap", "currency": "EUR", "lines": [dict(CASES_IN["subscription_growth"], name="A"), dict(CASES_IN["project"], name="B")]})
        with tempfile.TemporaryDirectory() as d6:
            _fake_charts(d6, mm, line="A")
            check("report: multi-line without 'line' rejected", br.check_requirements(_minimal_report(), mm, d6, T) != [])
            check("report: dashboard rendered for line A rejected for a report on line B", any("dashboard.json was rendered for line" in e for e in br.check_charts(d6, mm, {"line": "B", "lang": "en", "currency": "$"})))
            check("report: dashboard rendered in another language / currency rejected", any("lang=" in e for e in br.check_charts(d6, mm, {"line": "A", "lang": "ru", "currency": "$"})) and any("currency=" in e for e in br.check_charts(d6, mm, {"line": "A", "lang": "en", "currency": "€"})))
            stories = [dict(good["assumptions"]["stories"][0], domain=d, line=ln, stakes=f"Capital need {{{{combined.market.peak_capital_need|money}}}} → {{{{combined.bridge_by_domain.{ln}.{d}.peak_capital_need.if_founder_right|money}}}} if the team is right about this one number.") for ln, d in br.top_bridge_domains(mm, {"line": "A"})]
            nx_m = [dict(good["next_90_days"][0], domain=d) for ln, d in br.top_bridge_domains(mm, {"line": "A"}, k=1)] + good["next_90_days"][:2]
            good_m = _minimal_report(line="A", evidence_basis="None observed.", assumptions={"lead": LONG, "stories": stories}, next_90_days=nx_m)
            out6 = os.path.join(d6, "m.docx")
            ok6b, msg6b = _build(br, good_m, d6, out6, mm); check("report: multi-line report builds", ok6b, msg6b)
            doc6 = Document(out6); sizes = [len(t.rows) for t in doc6.tables[:5]]
            check("report: two lines (5 + 6 gaps) → page-1 table has header + 5 rows", 6 in sizes, sizes)
            ok6, msg6 = br.verify_pagination(out6, T); check("report: pagination holds with two lines", ok6 is True, msg6)


# ── Excel ────────────────────────────────────────────────────

def excel_checks(recalc):
    from openpyxl import load_workbook
    import importlib
    bw = importlib.import_module("build_workbook")
    CASES_IN["cac_zero"] = with_assumption(dict(CASES_IN["subscription"], founder_time_cost_per_customer=500), cac=A(0, 0))
    CASES_IN["cac_zero_oneoff"] = with_assumption(CASES_IN["one_off_ok"], cac=A(0, 0))
    for name in ("subscription", "subscription_growth", "project", "transactional", "one_off", "one_off_ok", "paid", "churn99", "trial_upfront",
                 "profitable_upfront", "post_revenue", "oneoff_cap", "horizon12", "channels", "cac_zero", "cac_zero_oneoff"):
        cfg = CASES_IN[name]
        with tempfile.TemporaryDirectory() as d:
            path, cases = bw.build_one(cfg, d, name)
            out = subprocess.run([sys.executable, recalc, path, "180"], capture_output=True, text=True)
            try:
                st = json.loads(out.stdout)
            except Exception:
                check(f"[{name}] recalc ran", False, out.stdout[-200:] + out.stderr[-200:]); continue
            check(f"[{name}] recalc 0 formula errors", st.get("total_errors") == 0, st.get("error_summary"))
            ws = load_workbook(path, data_only=True)["Summary"]
            rows = {r[0]: r[1:] for r in ws.iter_rows(min_row=8, max_row=26, max_col=5, values_only=True) if r[0]}
            n = cases["market"]["months"]
            for i, c in enumerate(CASES):
                py = cases[c]
                pairs = [("Blended CAC", py["inputs"]["blended_cac"]), ("LTV (lifetime gross-margin contribution, 60-mo cap)", py["unit_economics"]["ltv"]),
                         ("LTV / CAC (cash)", py["unit_economics"]["ltv_cac"] if py["unit_economics"]["ltv_cac"] is not None else "n/a"), ("Payback (months, cumulative)", py["unit_economics"]["payback_months"] or "none"),
                         ("Revenue month 12", py["milestones"]["m12"]["mrr"]), (f"Revenue month {n}", py["milestones"]["end"]["mrr"]),
                         ("Peak capital need", py["peak_capital_need"]), ("Break-even month", py["breakeven_month"] or "none")]
                for label, pv in pairs:
                    xv = rows[label][i] if label in rows else None
                    tol = (0.02, 0.0) if not isinstance(pv, (int, float)) or label not in ("Payback (months, cumulative)", "Break-even month") else (0.0, 0.5)
                    check(f"[{name}/{c}] {label}: xlsx {xv} ≈ py {pv}", close(xv, pv, *tol) if isinstance(pv, (int, float)) else close(xv, pv), "")
    for tag, cfg in [("multi", {"funding_path": "seed", "currency": "EUR", "lines": [dict(CASES_IN["subscription_growth"], name="Placement A", upfront_investment=5000), dict(CASES_IN["one_off_ok"], name="Workshops", upfront_investment=7000)], "shared_fixed_monthly": 300}),
                     ("collide", {"funding_path": "seed", "currency": "EUR", "lines": [dict(CASES_IN["subscription_growth"], name="Subscription basic"), dict(CASES_IN["one_off_ok"], name="Subscription pro", upfront_investment=7000)], "shared_fixed_monthly": 100})]:
        with tempfile.TemporaryDirectory() as d:
            path = bw.build_multi(cfg, d, tag)
            names = load_workbook(path).sheetnames
            check(f"[{tag}] 15 sheets, distinct case-insensitively, none auto-suffixed", len(names) == 15 and len({x.lower() for x in names}) == 15 and not any(x.endswith("1") for x in names), names)
            out = subprocess.run([sys.executable, recalc, path, "180"], capture_output=True, text=True)
            try:
                st = json.loads(out.stdout); check(f"[{tag}] recalc 0 formula errors", st.get("total_errors") == 0, st.get("error_summary"))
            except Exception:
                check(f"[{tag}] recalc ran", False, out.stdout[-200:] + out.stderr[-200:]); continue
            comb = run_all(cfg)["combined"]; ws = load_workbook(path, data_only=True)["Combined"]
            for i, c in enumerate(CASES):
                be, cap, rev, cum = [ws.cell(row=9 + i, column=k).value for k in (2, 3, 4, 5)]
                check(f"[{tag}/{c}] Combined break-even {be} ≈ py {comb[c]['breakeven_month']}", close(be, comb[c]["breakeven_month"] or "none"))
                check(f"[{tag}/{c}] Combined peak capital {cap} ≈ py {comb[c]['peak_capital_need']}", close(cap, comb[c]["peak_capital_need"]))
                check(f"[{tag}/{c}] Combined revenue last month ≈ py", close(rev, comb[c]["projection"][-1]["mrr"]))
                check(f"[{tag}/{c}] Combined cumulative cash ≈ py", close(cum, comb[c]["cumulative_m_last"]))
            check(f"[{tag}] Combined upfront ≈ py", close(ws["B7"].value, comb["market"]["upfront_investment_total"]), ws["B7"].value)
    # linked fills: a founder value filled from the market must follow the market range after a manual edit; unverified market follows the founder
    with tempfile.TemporaryDirectory() as d:
        path, cases = bw.build_one(CASES_IN["filled"], d, "filled")
        wb = load_workbook(path); ws = wb["Inputs"]
        rows = {ws.cell(row=r_, column=1).value: r_ for r_ in range(5, 20) if ws.cell(row=r_, column=1).value}
        pr, cr = rows["Price"], rows["Cost to acquire"]
        check("[filled] founder price cell is a formula over the market range", str(ws.cell(row=pr, column=2).value).startswith("=AVERAGE"), ws.cell(row=pr, column=2).value)
        check("[filled] unverified market CAC cells are formulas over the founder cell", str(ws.cell(row=cr, column=4).value) == f"=B{cr}" and str(ws.cell(row=cr, column=5).value) == f"=B{cr}")
        ws.cell(row=pr, column=4).value = 100; ws.cell(row=pr, column=5).value = 120; ws.cell(row=cr, column=2).value = 200; wb.save(path)
        subprocess.run([sys.executable, recalc, path, "180"], capture_output=True, text=True)
        wv = load_workbook(path, data_only=True)
        check("[filled] after editing the market range 100–120, the founder price follows (110)", close(wv["Inputs"].cell(row=pr, column=2).value, 110), wv["Inputs"].cell(row=pr, column=2).value)
        check("[filled] after editing the founder CAC to 200, the market CAC follows", close(wv["Inputs"].cell(row=cr, column=3).value, 200), wv["Inputs"].cell(row=cr, column=3).value)
        check("[filled] Founder sheet uses the edited values (ARPU 110, CAC 200)", close(wv["Founder"]["B3"].value, 110) and close(wv["Founder"]["B7"].value, 200), (wv["Founder"]["B3"].value, wv["Founder"]["B7"].value))
    with tempfile.TemporaryDirectory() as d:
        path, cases = bw.build_one(CASES_IN["channels"], d, "ch")
        ws = load_workbook(path)["Inputs"]; cr = next(r_ for r_ in range(5, 20) if ws.cell(row=r_, column=1).value == "Cost to acquire")
        check("[channels] founder CAC on Inputs is a formula over the channel table", "SUMPRODUCT" in str(ws.cell(row=cr, column=2).value), ws.cell(row=cr, column=2).value)
    with tempfile.TemporaryDirectory() as d:
        cfg_d = {**CASES_IN["channels"], "assumptions": {k: v for k, v in CASES_IN["channels"]["assumptions"].items() if k != "cac"}}   # CAC derived from channels on both sides
        path, cases = bw.build_one(cfg_d, d, "chd")
        wb = load_workbook(path); ws = wb["Inputs"]; cr = next(r_ for r_ in range(5, 20) if ws.cell(row=r_, column=1).value == "Cost to acquire")
        cs = next(r_ for r_ in range(5, 40) if ws.cell(row=r_, column=1).value == "a"); ws.cell(row=cs, column=3).value = 50; wb.save(path)
        subprocess.run([sys.executable, recalc, path, "180"], capture_output=True, text=True)
        wv = load_workbook(path, data_only=True)
        cfg_e = json.loads(json.dumps(cfg_d)); cfg_e["channels"][0]["cac"] = 50; py = run_line(cfg_e)["cases"]
        check("[channels-derived] after editing a channel CAC, founder AND market cases follow (Excel == Python)",
              close(wv["Founder"]["B7"].value, py["founder"]["inputs"]["blended_cac"]) and close(wv["Market"]["B7"].value, py["market"]["inputs"]["blended_cac"]), (wv["Founder"]["B7"].value, wv["Market"]["B7"].value, py["market"]["inputs"]["blended_cac"]))
    # the bundled recalculator must actually recalculate, including a workbook whose cached values are stale
    with tempfile.TemporaryDirectory() as d:
        from openpyxl import Workbook
        wb1 = Workbook(); w1 = wb1.active; w1["A1"] = 2; w1["A2"] = "=A1*21"; p1 = os.path.join(d, "one.xlsx"); wb1.save(p1)
        out = subprocess.run([sys.executable, os.path.join(HERE, "recalc.py"), p1, "180"], capture_output=True, text=True)
        try: st = json.loads(out.stdout)
        except Exception: st = {"error": out.stdout[-200:] + out.stderr[-200:]}
        v = load_workbook(p1, data_only=True).active["A2"].value if st.get("status") == "success" else None
        check("[recalc] bundled recalc.py recalculates a one-formula workbook (A2 = 42)", v == 42, st)
        # patch A1 to 3 in the sheet XML, keeping the cached <v>42</v> — exactly the case LibreOffice would not recompute on its own
        import zipfile, re as _re
        with zipfile.ZipFile(p1) as z: names = z.namelist(); data = {nme: z.read(nme) for nme in names}
        sh = [nme for nme in names if nme.startswith("xl/worksheets/sheet")][0]
        data[sh] = _re.sub(r'(<c r="A1"[^>]*>\s*<v>)2(</v>)', r'\g<1>3\2', data[sh].decode()).encode()
        with zipfile.ZipFile(p1, "w", zipfile.ZIP_DEFLATED) as zo:
            for nme in names: zo.writestr(nme, data[nme])
        check("[recalc] stale fixture: A1 = 3 with cached A2 = 42", load_workbook(p1, data_only=True).active["A1"].value == 3 and load_workbook(p1, data_only=True).active["A2"].value == 42)
        out = subprocess.run([sys.executable, os.path.join(HERE, "recalc.py"), p1, "180"], capture_output=True, text=True)
        v2 = load_workbook(p1, data_only=True).active["A2"].value
        check("[recalc] stale cached value is recomputed after an input edit (A2 = 63)", v2 == 63, (v2, out.stdout[-120:]))


if __name__ == "__main__":
    python_checks()
    recalc = sys.argv[sys.argv.index("--recalc") + 1] if "--recalc" in sys.argv else (os.path.join(HERE, "recalc.py") if shutil.which("soffice") else None)
    excel_ran = False
    if "--python-only" in sys.argv:
        print("WARNING: Excel checks skipped (--python-only). Not a release-grade run.")
    elif recalc:
        excel_checks(recalc); excel_ran = True
    else:
        print("INCOMPLETE: Excel ↔ Python comparison did not run (LibreOffice `soffice` not on PATH). Install LibreOffice or pass --recalc <path>. Do not deliver from this build.")
        sys.exit(1)
    if FAILS:
        print(f"\n{len(FAILS)} failures"); sys.exit(1)
    print("\nALL PASS (python + excel)" if excel_ran else "\nPYTHON PASS — Excel not verified")
    sys.exit(0)
